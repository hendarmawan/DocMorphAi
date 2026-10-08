"""Builds realistic test documents programmatically, so fixtures stay reviewable."""

from __future__ import annotations

import base64
import io

from docx import Document
from docx.opc.constants import RELATIONSHIP_TYPE as RT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches

# 1x1 transparent PNG
PNG_1PX = base64.b64decode(
    "iVBORw0KGgoAAAANSUhEUgAAAAEAAAABCAYAAAAfFcSJAAAADUlEQVR42mNkYPhfDwAChwGA60e6kgAAAABJRU5ErkJggg=="
)


def _add_hyperlink(paragraph, text: str, url: str) -> None:
    part = paragraph.part
    r_id = part.relate_to(url, RT.HYPERLINK, is_external=True)
    link = OxmlElement("w:hyperlink")
    link.set(qn("r:id"), r_id)
    run = OxmlElement("w:r")
    t = OxmlElement("w:t")
    t.text = text
    run.append(t)
    link.append(run)
    paragraph._p.append(link)


def make_docx(*, link_url: str = "https://example.com/report", with_image: bool = True) -> bytes:
    doc = Document()
    doc.core_properties.title = "Quarterly Strategy Report"
    doc.core_properties.author = "DocMorph QA"
    doc.add_heading("Quarterly Strategy Report", level=0)
    doc.add_heading("Executive summary", level=1)
    p = doc.add_paragraph("Revenue grew ")
    p.add_run("18%").bold = True
    p.add_run(" this quarter, driven by ")
    p.add_run("enterprise").italic = True
    p.add_run(" demand. See ")
    _add_hyperlink(p, "the full report", link_url)
    p.add_run(".")
    doc.add_heading("Priorities", level=2)
    doc.add_paragraph("Expand the partner market", style="List Bullet")
    doc.add_paragraph("Reduce onboarding time", style="List Bullet")
    doc.add_paragraph("Hire two engineers", style="List Number")
    doc.add_paragraph("Ship release one", style="List Number")
    doc.add_heading("Budget", level=1)
    table = doc.add_table(rows=3, cols=2)
    for r, (a, b) in enumerate([("Item", "Amount"), ("Cloud", "$12,000"), ("Staff", "$80,000")]):
        table.cell(r, 0).text = a
        table.cell(r, 1).text = b
    for cell in table.rows[0].cells:
        for run in cell.paragraphs[0].runs:
            run.bold = True
    doc.add_paragraph("Budget stakeholder sign-off is due next quarter.", style="Quote")
    if with_image:
        doc.add_heading("Appendix", level=1)
        doc.add_picture(io.BytesIO(PNG_1PX), width=Inches(1))
    buf = io.BytesIO()
    doc.save(buf)
    return buf.getvalue()


def make_pdf(text_lines: list[str]) -> bytes:
    """Minimal single-page PDF with a real text layer (no external tools needed)."""
    content_lines = ["BT", "/F1 12 Tf", "72 720 Td", "14 TL"]
    for line in text_lines:
        safe = line.replace("\\", "\\\\").replace("(", "\\(").replace(")", "\\)")
        content_lines.append(f"({safe}) Tj T*")
    content_lines.append("ET")
    stream = "\n".join(content_lines).encode("latin-1")
    objects = [
        b"<< /Type /Catalog /Pages 2 0 R >>",
        b"<< /Type /Pages /Kids [3 0 R] /Count 1 >>",
        b"<< /Type /Page /Parent 2 0 R /MediaBox [0 0 612 792] /Contents 4 0 R "
        b"/Resources << /Font << /F1 5 0 R >> >> >>",
        b"<< /Length " + str(len(stream)).encode() + b" >>\nstream\n" + stream + b"\nendstream",
        b"<< /Type /Font /Subtype /Type1 /BaseFont /Helvetica >>",
    ]
    out = io.BytesIO()
    out.write(b"%PDF-1.4\n")
    offsets = []
    for i, obj in enumerate(objects, start=1):
        offsets.append(out.tell())
        out.write(f"{i} 0 obj\n".encode() + obj + b"\nendobj\n")
    xref = out.tell()
    out.write(f"xref\n0 {len(objects) + 1}\n0000000000 65535 f \n".encode())
    for off in offsets:
        out.write(f"{off:010d} 00000 n \n".encode())
    out.write(f"trailer\n<< /Size {len(objects) + 1} /Root 1 0 R >>\nstartxref\n{xref}\n%%EOF\n".encode())
    return out.getvalue()


SAMPLE_MARKDOWN = """# Research Notes

## Abstract

This study examines **adaptive** reading and the *methodology* behind it.

## Results

1. First finding
2. Second finding

| Metric | Value |
|--------|-------|
| Speed  | 1.4x  |

> Reading is a conversation.

```python
print("hello")
```

---

See [the paper](https://example.org/paper) for references.
"""
